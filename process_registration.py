"""
Excel and Word File Processing Automation
Run: python process_registration.py
Required packages: pandas, openpyxl, python-docx
"""
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "registration_raw.xlsx"
CLEAN_FILE = BASE_DIR / "registration_cleaned.xlsx"
WORD_FILE = BASE_DIR / "registration_report.docx"

def clean_registration_data(df):
    # Remove leading/trailing spaces from text columns
    for column in df.select_dtypes(include="object").columns:
        df[column] = df[column].map(lambda value: value.strip() if isinstance(value, str) else value)

    # Standardize email addresses and course names
    df["Email"] = df["Email"].map(lambda value: value.lower() if isinstance(value, str) else value)
    df["Course"] = df["Course"].str.title()
    df["City"] = df["City"].str.title()

    # Convert registration dates to a consistent date format; invalid values become blank
    df["Registration Date"] = pd.to_datetime(
        df["Registration Date"], errors="coerce"
    ).dt.strftime("%Y-%m-%d")

    # Replace empty strings with missing values, then fill missing values clearly
    df = df.replace(r"^\s*$", pd.NA, regex=True)
    df["Email"] = df["Email"].fillna("Not provided")
    df["Phone"] = df["Phone"].fillna("Not provided")
    df["Registration Date"] = df["Registration Date"].fillna("Not provided")
    df["City"] = df["City"].fillna("Not provided")

    # Remove duplicate student records based on Student ID, keeping the first record
    df = df.drop_duplicates(subset=["Student ID"], keep="first")
    return df.sort_values("Student ID").reset_index(drop=True)

def create_word_report(df, removed_duplicates, output_path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)

    title = doc.add_heading("Student Registration Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(
        "This report was generated automatically using Python from the cleaned Excel registration data."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(f"Unique student records: {len(df)}")
    doc.add_paragraph(f"Duplicate records removed: {removed_duplicates}")
    doc.add_paragraph(f"Missing or invalid values are labeled 'Not provided'.")

    doc.add_heading("Registration Details", level=1)
    columns = ["Student ID", "Full Name", "Email", "Course", "Registration Date", "City"]
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, col in enumerate(columns):
        table.rows[0].cells[idx].text = col
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for idx, col in enumerate(columns):
            value = row[col]
            cells[idx].text = "" if pd.isna(value) else str(value)

    doc.add_heading("Cleaning Steps Applied", level=1)
    for item in [
        "Trimmed unnecessary spaces from text fields.",
        "Converted email addresses to lowercase.",
        "Standardized course and city names.",
        "Converted dates to YYYY-MM-DD and flagged invalid dates.",
        "Handled missing values using the label 'Not provided'.",
        "Removed duplicate records using Student ID.",
        "Sorted records by Student ID.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(output_path)

def main():
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Input file not found: {INPUT_FILE}. Keep registration_raw.xlsx beside this script."
        )

    raw_df = pd.read_excel(INPUT_FILE)
    before_count = len(raw_df)
    cleaned_df = clean_registration_data(raw_df.copy())
    removed_duplicates = before_count - len(cleaned_df)

    cleaned_df.to_excel(CLEAN_FILE, index=False, sheet_name="Cleaned Registrations")
    create_word_report(cleaned_df, removed_duplicates, WORD_FILE)

    print("Excel and Word processing completed successfully!")
    print(f"Original records: {before_count}")
    print(f"Cleaned records: {len(cleaned_df)}")
    print(f"Duplicate records removed: {removed_duplicates}")
    print(f"Saved: {CLEAN_FILE.name}")
    print(f"Saved: {WORD_FILE.name}")

if __name__ == "__main__":
    main()
